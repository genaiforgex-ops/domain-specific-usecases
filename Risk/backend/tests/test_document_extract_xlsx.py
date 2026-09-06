"""Unit tests for the .xlsx branch of document_extract.extract_text().

Covers the grid-to-text shape the classifier is fed: one line per row, cells
joined with " | ", merged ranges collapsed, trailing empty columns trimmed, the
separator/newline sanitising that keeps one row on one line, visible-sheet
selection, cell-value coercion, and the failure modes (unreadable workbook,
unsupported extension).

Does NOT cover extract_canonical_fields / summarize_selected_answers semantics —
unchanged by this work, and the classification path deliberately routes .xlsx
around the former (see app/api/v1/classifications.py). The upload endpoint itself
writes to the DB and the repo has no DB fixtures, so it is covered by the manual
verification steps instead.

Note there is no highlight/fill test here by design: in these workbooks the
question sits in one column and the answer in another, so no answer is carried by
cell formatting. Ticked-glyph marking is inherited from extract_text's shared
_mark_checked_boxes pass and is asserted once below.
"""
import io
from datetime import date, datetime
from pathlib import Path

import pytest
from openpyxl import Workbook

from app.services.document_extract import extract_text


def _book(sheets: dict[str, list[list]], *, hidden: tuple[str, ...] = ()) -> bytes:
    """{title: rows-of-cell-values} -> .xlsx bytes. First sheet keeps workbook order."""
    wb = Workbook()
    wb.remove(wb.active)
    for title, rows in sheets.items():
        ws = wb.create_sheet(title=title)
        for row in rows:
            ws.append(row)
        if title in hidden:
            ws.sheet_state = "hidden"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _merged_book(rows: list[list], merges: tuple[str, ...]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "S"
    for row in rows:
        ws.append(row)
    for rng in merges:
        ws.merge_cells(rng)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _extract(data: bytes) -> str:
    return extract_text("vdd.xlsx", data)


def _body(text: str) -> list[str]:
    """Lines excluding the sheet headings and their blank separators."""
    return [ln for ln in text.split("\n") if ln.strip() and not ln.startswith("=== Sheet:")]


# ---- shape ---------------------------------------------------------------

def test_an_xlsx_upload_is_extracted_as_text_instead_of_zip_mojibake():
    # The regression this branch exists for: an .xlsx is a zip container, so the
    # old utf-8 fallback returned the PK header and the rest of the archive.
    text = _extract(_book({"S": [["Name of the service provider", "Acme Ltd"]]}))

    assert not text.startswith("PK")
    assert "Name of the service provider | Acme Ltd" in text


def test_each_row_becomes_one_line_with_cells_joined_by_a_pipe():
    text = _extract(_book({"S": [["1", "Parameter", "Response"], ["2", "Other", "Answer"]]}))

    assert _body(text) == ["1 | Parameter | Response", "2 | Other | Answer"]


def test_a_row_with_nothing_in_it_is_skipped():
    text = _extract(_book({"S": [["a"], [None, None], ["b"]]}))

    assert _body(text) == ["a", "b"]


def test_trailing_empty_columns_do_not_pad_the_line():
    # Excel rows are as wide as the sheet's used range, not the table.
    text = _extract(_book({"S": [["a", "b", None, None], ["c", "d", "e", "f"]]}))

    assert _body(text)[0] == "a | b"


def test_an_empty_cell_between_two_filled_ones_is_kept_so_columns_stay_aligned():
    text = _extract(_book({"S": [["Question", None, "Answer"]]}))

    assert _body(text) == ["Question |  | Answer"]


def test_a_merged_parameter_span_emits_its_value_once_and_skips_the_covered_columns():
    # A real VDD sheet merges Parameter across B:D; without collapsing, the line
    # reads "1 | Parameter |  |  | Response".
    text = _extract(_merged_book([["1", "Parameter", None, None, "Response"]], ("B1:D1",)))

    assert _body(text) == ["1 | Parameter | Response"]


def test_a_vertically_merged_cell_repeats_nothing_on_the_rows_below_it():
    text = _extract(_merged_book([["Section", "one"], [None, "two"]], ("A1:A2",)))

    assert _body(text) == ["Section | one", "two"]


# ---- value coercion (pins the shared cell_to_str) -------------------------

def test_a_date_cell_reads_as_an_iso_date_not_a_midnight_timestamp():
    text = _extract(_book({"S": [["Commencement", datetime(2026, 8, 19)]]}))

    assert "Commencement | 2026-08-19" in text
    assert "00:00" not in text


def test_a_plain_date_cell_reads_as_an_iso_date():
    text = _extract(_book({"S": [["Commencement", date(2026, 8, 19)]]}))

    assert "Commencement | 2026-08-19" in text


def test_a_whole_number_cell_has_no_decimal_tail():
    text = _extract(_book({"S": [["Years in operation", 12.0]]}))

    assert "Years in operation | 12" in text
    assert "12.0" not in text


def test_a_boolean_cell_reads_as_yes_or_no():
    text = _extract(_book({"S": [["Has DR", True], ["Tested", False]]}))

    assert "Has DR | Yes" in text
    assert "Tested | No" in text


# ---- separator safety ----------------------------------------------------

def test_a_pipe_typed_inside_a_cell_cannot_fake_a_column_break():
    text = _extract(_book({"S": [["Data types", "PII | Financial"]]}))

    assert _body(text) == ["Data types | PII / Financial"]


def test_a_wrapped_cell_with_newlines_stays_on_one_line():
    text = _extract(_book({"S": [["Scope", "first line\nsecond line"]]}))

    assert _body(text) == ["Scope | first line second line"]


# ---- sheets --------------------------------------------------------------

def test_every_visible_sheet_is_extracted_not_just_the_first():
    text = _extract(_book({"First": [["a"]], "Second": [["b"]]}))

    assert "a" in text
    assert "b" in text


def test_a_hidden_sheet_is_skipped():
    # Hidden sheets in a VDD workbook are scoring rubrics, not the form.
    text = _extract(_book({"Form": [["real"]], "Rubric": [["noise"]]}, hidden=("Rubric",)))

    assert "real" in text
    assert "noise" not in text


def test_each_sheet_is_introduced_by_a_heading_line():
    text = _extract(_book({"Vendor Exit Strategy": [["a"]]}))

    assert "=== Sheet: Vendor Exit Strategy ===" in text


def test_a_blank_sheet_contributes_no_heading():
    text = _extract(_book({"Full": [["a"]], "Empty": [[None]]}))

    assert "=== Sheet: Empty ===" not in text
    assert "=== Sheet: Full ===" in text


# ---- glyphs (inherited from the shared post-pass) ------------------------

def test_a_checkbox_glyph_in_a_cell_still_becomes_a_selected_marker():
    text = _extract(_book({"S": [["IT outsourcing", "Yes ☐ No ☑"]]}))

    assert "[SELECTED: No ☑]" in text


# ---- failure modes -------------------------------------------------------

def test_a_formula_cell_with_no_cached_value_reads_as_blank_rather_than_the_formula():
    # data_only=True returns the cached result, and a workbook never opened by
    # Excel has no cache. Blank is honest; "=B2*C2" in the prompt would not be.
    text = _extract(_book({"S": [["Total", "=1+1"]]}))

    assert "=1+1" not in text
    assert "Total" in text


def test_a_file_that_is_not_a_workbook_yields_empty_text_rather_than_binary_noise():
    assert extract_text("x.xlsx", b"this is not a spreadsheet") == ""


def test_an_unsupported_extension_is_rejected_instead_of_decoded_as_text():
    with pytest.raises(ValueError, match="Unsupported file type"):
        extract_text("scan.csv", b"\x00\x01binary")


def test_the_still_supported_types_are_unaffected_by_the_new_guard():
    assert extract_text("notes.txt", b"hello") == "hello"


# ---- the real VDD workbook ------------------------------------------------
# Guarded the same way test_form_import.py guards it: the file lives at the repo
# root, not in the test tree, so a checkout without it skips rather than fails.

WORKBOOK = Path(__file__).resolve().parents[2] / "1 Vendor Due Diligence - [Vendor Name].xlsx"


@pytest.fixture(scope="module")
def real_workbook_text() -> str:
    if not WORKBOOK.exists():
        pytest.skip("VDD workbook not present")
    return extract_text(WORKBOOK.name, WORKBOOK.read_bytes())


def test_the_real_vdd_workbook_yields_the_questions_it_asks(real_workbook_text):
    assert not real_workbook_text.startswith("PK")
    assert "Name of the service provider" in real_workbook_text
    assert "Sl. No. | Parameter | Type | Response" in real_workbook_text


def test_the_real_vdd_workbook_leaves_out_its_hidden_scoring_sheets(real_workbook_text):
    # These strings appear only on hidden rubric sheets. If they show up, the
    # visible-sheet filter has regressed and the classifier is being fed the
    # scoring methodology alongside the form.
    assert "Inherent risk scoring" not in real_workbook_text
    assert "Net Outsourcing Risk Rating" not in real_workbook_text


def test_the_real_vdd_workbook_extracts_no_runaway_padding(real_workbook_text):
    # The sheets declare a used range far wider and taller than their content;
    # without the trailing-cell trim and the blank-row skip this balloons.
    assert 5_000 < len(real_workbook_text) < 15_000
    assert " |  |  |  | " not in real_workbook_text
