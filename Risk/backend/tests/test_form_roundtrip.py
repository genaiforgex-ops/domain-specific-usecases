"""Unit tests for the form Excel round-trip.

Covers the parts that don't need a database: workbook write -> read fidelity
across every field type, the cell-type coercion that keeps openpyxl's native
date/number/bool values valid against the API schema, draft merge semantics,
and the label-matching fallback. The two endpoints wrapping this also enforce
ownership and template-mismatch rules; those write to the DB and the repo has
no DB fixtures, so they are covered by the manual verification steps instead.
"""
import io
from datetime import date, datetime

import pytest
from openpyxl import load_workbook

from app.schemas.form import FormAssignmentDraftUpdate
from app.services.form_export import (
    COL_FIELD_ID,
    COL_GUIDANCE,
    COL_RESPONSE,
    MAIN_SHEET,
    META_SHEET,
    build_assignment_workbook,
    find_header_row,
    merge_answers,
    parse_answer_workbook,
)

SCHEMA = {
    "sections": [
        {
            "id": "identity",
            "title": "Entity Identity",
            "fields": [
                {"id": "q1_name", "label": "Name of the Service Provider", "type": "text", "required": True},
                {"id": "q2_desc", "label": "Describe the service", "type": "textarea", "required": False},
                {
                    "id": "q3_encrypted",
                    "label": "Is data encrypted at rest?",
                    "type": "yes_no",
                    "options": ["Yes", "No"],
                    "required": True,
                },
            ],
        },
        {
            "id": "commercial",
            "title": "Commercial",
            "fields": [
                {"id": "q4_headcount", "label": "Employee headcount", "type": "number", "required": False},
                {"id": "q5_start", "label": "Contract start date", "type": "date", "required": False},
                {
                    "id": "q6_tier",
                    "label": "Criticality tier",
                    "type": "select",
                    "options": ["Critical", "Material", "Low"],
                    "required": False,
                },
            ],
        },
    ],
    "repeatable_groups": [
        {
            "id": "subcontractors",
            "title": "Subcontractors",
            "fields": [
                {"id": "sub_name", "label": "Subcontractor name", "type": "text"},
                {"id": "sub_country", "label": "Country", "type": "text"},
            ],
        }
    ],
}

ANSWERS = {
    "q1_name": "Acme IT Services Pvt Ltd",
    "q2_desc": "Managed hosting and\napplication support",
    "q3_encrypted": "Yes",
    "q4_headcount": "420",
    "q5_start": "2026-04-01",
    "q6_tier": "Material",
    "subcontractors": [
        {"sub_name": "Globex Ltd", "sub_country": "IN"},
        {"sub_name": "Initech Inc", "sub_country": "US"},
    ],
}

META = {
    "title": "Vendor Due Diligence",
    "vendor_legal_name": "Acme IT Services Pvt Ltd",
    "assignment_id": "11111111-1111-1111-1111-111111111111",
    "template_id": "22222222-2222-2222-2222-222222222222",
    "template_version": 3,
    "exported_at": "2026-08-13T10:00:00+00:00",
}


def _rewrite_responses(data: bytes, updates: dict[str, object]) -> bytes:
    """Edit the Response cells of an exported workbook, the way a vendor would.

    Values are written with their native Python type on purpose — that is what
    Excel produces when someone types a date or a number, and it is exactly the
    case the string coercion has to survive.
    """
    wb = load_workbook(io.BytesIO(data))
    ws = wb[MAIN_SHEET]
    for row in ws.iter_rows(min_row=find_header_row(ws) + 1, max_col=COL_FIELD_ID):
        field_id = row[COL_FIELD_ID - 1].value
        if field_id in updates:
            row[COL_RESPONSE - 1].value = updates[field_id]
    buf = io.BytesIO()
    wb.save(buf)
    wb.close()
    return buf.getvalue()


def test_roundtrip_preserves_every_field_type():
    data = build_assignment_workbook(SCHEMA, ANSWERS, META)
    parsed = parse_answer_workbook(data, SCHEMA)

    assert parsed.answers == ANSWERS
    assert parsed.unmatched == []
    assert parsed.warnings == []


def test_roundtrip_carries_metadata():
    parsed = parse_answer_workbook(build_assignment_workbook(SCHEMA, ANSWERS, META), SCHEMA)

    assert parsed.meta["template_id"] == META["template_id"]
    assert parsed.meta["assignment_id"] == META["assignment_id"]
    assert parsed.meta["template_version"] == "3"


def test_blank_export_has_no_answers_but_keeps_the_questions():
    parsed = parse_answer_workbook(build_assignment_workbook(SCHEMA, {}, META), SCHEMA)
    assert parsed.answers == {}

    wb = load_workbook(io.BytesIO(build_assignment_workbook(SCHEMA, {}, META)))
    ws = wb[MAIN_SHEET]
    labels = [c.value for c in ws["B"][find_header_row(ws):] if c.value]
    assert "Name of the Service Provider" in labels
    assert "Criticality tier" in labels


@pytest.mark.parametrize(
    "typed_value,expected",
    [
        (datetime(2026, 4, 1, 0, 0), "2026-04-01"),
        (date(2026, 4, 1), "2026-04-01"),
        (420, "420"),
        (420.0, "420"),
        (99.5, "99.5"),
        (True, "Yes"),
        (False, "No"),
    ],
)
def test_native_excel_cell_types_coerce_to_strings(typed_value, expected):
    """A natively-typed cell must not reach the API as a date/int/bool.

    FormAssignmentDraftUpdate.answers is dict[str, str | list[dict[str, str]]],
    so anything else would 422 on save.
    """
    data = _rewrite_responses(
        build_assignment_workbook(SCHEMA, {}, META), {"q4_headcount": typed_value}
    )
    parsed = parse_answer_workbook(data, SCHEMA)

    assert parsed.answers["q4_headcount"] == expected
    FormAssignmentDraftUpdate(answers=parsed.answers)


def test_parsed_answers_validate_against_the_draft_schema():
    parsed = parse_answer_workbook(build_assignment_workbook(SCHEMA, ANSWERS, META), SCHEMA)
    validated = FormAssignmentDraftUpdate(answers=parsed.answers)

    assert validated.answers["q1_name"] == "Acme IT Services Pvt Ltd"
    assert validated.answers["subcontractors"][1]["sub_country"] == "US"


def test_free_text_over_a_dropdown_normalises_to_the_canonical_option():
    data = _rewrite_responses(
        build_assignment_workbook(SCHEMA, {}, META),
        {"q3_encrypted": "yes", "q6_tier": "MATERIAL"},
    )
    parsed = parse_answer_workbook(data, SCHEMA)

    assert parsed.answers["q3_encrypted"] == "Yes"
    assert parsed.answers["q6_tier"] == "Material"


def test_blank_cells_are_dropped_rather_than_imported_as_empty():
    data = _rewrite_responses(
        build_assignment_workbook(SCHEMA, ANSWERS, META), {"q2_desc": None, "q6_tier": "   "}
    )
    parsed = parse_answer_workbook(data, SCHEMA)

    assert "q2_desc" not in parsed.answers
    assert "q6_tier" not in parsed.answers


def test_merge_lets_the_imported_file_win_without_wiping_untouched_answers():
    draft = {"q1_name": "Acme", "q2_desc": "Old notes", "q6_tier": "Low"}
    incoming = {"q1_name": "Acme IT Services Pvt Ltd", "q3_encrypted": "Yes"}

    merged = merge_answers(draft, incoming)

    assert merged["q1_name"] == "Acme IT Services Pvt Ltd"  # imported wins
    assert merged["q3_encrypted"] == "Yes"                  # imported adds
    assert merged["q2_desc"] == "Old notes"                 # untouched survives
    assert merged["q6_tier"] == "Low"
    assert draft["q1_name"] == "Acme"                       # caller's dict untouched


def test_merge_handles_a_draft_that_was_never_saved():
    assert merge_answers(None, {"q1_name": "Acme"}) == {"q1_name": "Acme"}


def test_answers_still_match_when_the_hidden_id_column_is_stripped():
    """A vendor who copies the rows into a fresh workbook loses column F."""
    wb = load_workbook(io.BytesIO(build_assignment_workbook(SCHEMA, ANSWERS, META)))
    ws = wb[MAIN_SHEET]
    for row in ws.iter_rows(min_row=1, max_col=COL_FIELD_ID):
        row[COL_FIELD_ID - 1].value = None
    del wb[META_SHEET]
    buf = io.BytesIO()
    wb.save(buf)
    wb.close()

    parsed = parse_answer_workbook(buf.getvalue(), SCHEMA)

    assert parsed.answers["q1_name"] == "Acme IT Services Pvt Ltd"
    assert parsed.answers["q3_encrypted"] == "Yes"
    assert parsed.meta == {}
    assert any("no GenAIForge Risk metadata" in w for w in parsed.warnings)


def test_rows_that_match_nothing_are_reported_not_silently_dropped():
    wb = load_workbook(io.BytesIO(build_assignment_workbook(SCHEMA, {}, META)))
    ws = wb[MAIN_SHEET]
    row = ws.max_row + 1
    ws.cell(row=row, column=2, value="A question this template does not have")
    ws.cell(row=row, column=COL_RESPONSE, value="Some answer")
    buf = io.BytesIO()
    wb.save(buf)
    wb.close()

    parsed = parse_answer_workbook(buf.getvalue(), SCHEMA)

    assert parsed.answers == {}
    assert parsed.unmatched == ["A question this template does not have"]


def test_a_template_without_repeatable_groups_exports_cleanly():
    """form_import.py always emits repeatable_groups: [], so this is the norm."""
    schema = {"sections": SCHEMA["sections"], "repeatable_groups": []}
    parsed = parse_answer_workbook(build_assignment_workbook(schema, ANSWERS, META), schema)

    assert parsed.answers["q1_name"] == "Acme IT Services Pvt Ltd"
    assert "subcontractors" not in parsed.answers


def _guidance_column(data: bytes) -> dict[str, str]:
    """field_id -> guidance cell, as an offline filler would read it."""
    wb = load_workbook(io.BytesIO(data))
    ws = wb[MAIN_SHEET]
    out = {}
    for row in ws.iter_rows(min_row=find_header_row(ws) + 1, max_col=COL_FIELD_ID):
        fid = row[COL_FIELD_ID - 1].value
        if fid:
            out[fid] = row[COL_GUIDANCE - 1].value or ""
    wb.close()
    return out


def test_help_text_travels_into_the_exported_workbook():
    """Whoever fills the file offline needs the same guidance the wizard shows."""
    schema = {
        "sections": [
            {
                "id": "s1",
                "title": "Scope",
                "fields": [
                    {
                        "id": "q1_scope",
                        "label": "Scope of the activity",
                        "type": "textarea",
                        "help_text": "Below Example for Reference:\nFront-line call centre operations.",
                    }
                ],
            }
        ],
        "repeatable_groups": [],
    }
    assert _guidance_column(build_assignment_workbook(schema, {}, META))["q1_scope"].startswith(
        "Below Example for Reference:"
    )


def test_guidance_is_not_read_back_as_an_answer():
    schema = {
        "sections": [
            {
                "id": "s1",
                "title": "Scope",
                "fields": [{"id": "q1_scope", "label": "Scope", "type": "text", "help_text": "Please provide detail"}],
            }
        ],
        "repeatable_groups": [],
    }
    parsed = parse_answer_workbook(build_assignment_workbook(schema, {}, META), schema)

    assert parsed.answers == {}


def test_a_form_description_becomes_a_banner_without_breaking_the_round_trip():
    """The banner shifts the header down a row; the reader must find it."""
    schema = {**SCHEMA, "description": "Answer every starred question. Attach policy names and dates."}
    data = build_assignment_workbook(schema, ANSWERS, META)

    wb = load_workbook(io.BytesIO(data))
    ws = wb[MAIN_SHEET]
    assert find_header_row(ws) == 3
    assert "Answer every starred question" in str(ws.cell(row=2, column=1).value)
    wb.close()

    assert parse_answer_workbook(data, schema).answers == ANSWERS


def test_unreadable_upload_raises_a_clean_error():
    with pytest.raises(ValueError, match="Could not read the workbook"):
        parse_answer_workbook(b"this is not a spreadsheet", SCHEMA)
