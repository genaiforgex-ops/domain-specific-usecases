"""Excel round-trip for form assignments.

The mirror image of form_import.py. That module turns a vendor's VDD workbook
*into* a template; this one turns a template + its answers back into a workbook
the vendor can fill offline, and reads that workbook back into answers.

The round-trip is keyed on a hidden `_field_id` column rather than on the
question text, so a vendor who rewords or reformats a label does not silently
orphan their answer. A hidden `_genaiforge` sheet carries the assignment/template
ids so an import can tell it is being given the right form. Matching by
normalised label is kept as a fallback for the case where someone copies the
rows into a fresh workbook and loses the id column.
"""

import io
import re
from dataclasses import dataclass, field as dataclass_field
from datetime import date, datetime, time
from typing import Any, Iterator

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill, Protection
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

FORMAT_VERSION = 1
MAIN_SHEET = "Due Diligence"
META_SHEET = "_genaiforge"

# Main sheet columns. Column F holds the field id and is hidden; it is what
# makes re-import exact rather than a guess. Guidance sits immediately right of
# Response so whoever is filling the sheet reads it where they type.
COL_SLNO, COL_PARAM, COL_TYPE, COL_RESPONSE, COL_GUIDANCE, COL_FIELD_ID = 1, 2, 3, 4, 5, 6
HEADERS = ["Sl. No.", "Parameter", "Section", "Response", "Guidance", "_field_id"]
# Where the header lands when the form carries no description. A description
# adds a banner row above it, so readers locate the header rather than assume.
DEFAULT_HEADER_ROW = 2

_HEADER_FILL = PatternFill("solid", fgColor="1F3A5F")
_SECTION_FILL = PatternFill("solid", fgColor="E8EDF3")
_INPUT_FILL = PatternFill("solid", fgColor="FFFDF0")


@dataclass
class ParsedAnswers:
    """Result of reading a filled workbook back in."""

    answers: dict[str, Any] = dataclass_field(default_factory=dict)
    meta: dict[str, str] = dataclass_field(default_factory=dict)
    unmatched: list[str] = dataclass_field(default_factory=list)
    warnings: list[str] = dataclass_field(default_factory=list)


def _normalise_label(text: str) -> str:
    return re.sub(r"[^a-z0-9]+", "", (text or "").lower())


def _safe_sheet_title(title: str, used: set[str]) -> str:
    """Excel sheet titles: max 31 chars, no []:*?/\\ , must be unique."""
    clean = re.sub(r"[\[\]:*?/\\]", " ", title or "Sheet").strip()[:31] or "Sheet"
    candidate = clean
    n = 2
    while candidate.lower() in used:
        suffix = f" {n}"
        candidate = clean[: 31 - len(suffix)] + suffix
        n += 1
    used.add(candidate.lower())
    return candidate


def cell_to_str(value: Any) -> str:
    """Coerce whatever openpyxl hands back into the plain string the API expects.

    `FormAssignmentDraftUpdate.answers` is typed dict[str, str | list[dict[str,
    str]]], so a natively-typed date or number cell would fail validation if it
    were passed through untouched.

    Public because document_extract's .xlsx branch needs the same coercion — a
    date cell must not reach the classifier as "2026-08-19 00:00:00".
    """
    if value is None:
        return ""
    if isinstance(value, bool):
        return "Yes" if value else "No"
    if isinstance(value, datetime):
        # A date-only cell still comes back as datetime at midnight.
        if value.time() == time(0, 0):
            return value.date().isoformat()
        return value.isoformat(sep=" ", timespec="minutes")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float):
        # Excel stores every number as a float; 42 must not become "42.0".
        if value.is_integer():
            return str(int(value))
        return repr(value).rstrip("0").rstrip(".")
    if isinstance(value, int):
        return str(value)
    return str(value).strip()


def _iter_section_fields(schema: dict) -> Iterator[tuple[dict, dict]]:
    """Yield (section, field) for the flat sections, in schema order."""
    for section in schema.get("sections") or []:
        for fld in section.get("fields") or []:
            yield section, fld


def _all_fields(schema: dict) -> Iterator[dict]:
    for _, fld in _iter_section_fields(schema):
        yield fld
    for group in schema.get("repeatable_groups") or []:
        for fld in group.get("fields") or []:
            yield fld


def _option_normaliser(schema: dict) -> dict[str, dict[str, str]]:
    """field_id -> {normalised option: canonical option}.

    Lets a vendor who typed "yes" over the dropdown still round-trip to "Yes",
    which is what the wizard's yes/no pills compare against.
    """
    out: dict[str, dict[str, str]] = {}
    for fld in _all_fields(schema):
        options = fld.get("options")
        if fld.get("type") == "yes_no" and not options:
            options = ["Yes", "No"]
        if options:
            out[fld["id"]] = {_normalise_label(o): o for o in options}
    return out


# --------------------------------------------------------------------------
# Write
# --------------------------------------------------------------------------


def build_assignment_workbook(
    schema: dict,
    answers: dict | None,
    meta: dict[str, Any] | None = None,
) -> bytes:
    """Render a template + answers as a fillable .xlsx."""
    answers = answers or {}
    meta = meta or {}

    wb = Workbook()
    ws = wb.active
    ws.title = MAIN_SHEET

    title = str(meta.get("title") or "Vendor Due Diligence")
    vendor = meta.get("vendor_legal_name")
    banner = f"{title} — {vendor}" if vendor else title
    ws.cell(row=1, column=COL_SLNO, value=banner).font = Font(bold=True, size=13)
    ws.merge_cells(start_row=1, start_column=COL_SLNO, end_row=1, end_column=COL_GUIDANCE)

    header_row = DEFAULT_HEADER_ROW
    description = (schema.get("description") or "").strip()
    if description:
        # The same "How to fill this form" text the web wizard shows, so someone
        # working only from the file is not left guessing.
        cell = ws.cell(row=2, column=COL_SLNO, value=description)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells(start_row=2, start_column=COL_SLNO, end_row=2, end_column=COL_GUIDANCE)
        ws.row_dimensions[2].height = min(90, 15 * (1 + len(description) // 90))
        header_row = 3

    for idx, header in enumerate(HEADERS, start=1):
        cell = ws.cell(row=header_row, column=idx, value=header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = _HEADER_FILL
        cell.alignment = Alignment(vertical="center")

    row = header_row + 1
    sl_no = 0
    last_section_id = None
    validations: dict[str, DataValidation] = {}

    for section, fld in _iter_section_fields(schema):
        if section.get("id") != last_section_id:
            last_section_id = section.get("id")
            heading = section.get("title") or ""
            if section.get("description"):
                heading = f"{heading} — {section['description']}"
            cell = ws.cell(row=row, column=COL_SLNO, value=heading)
            cell.font = Font(bold=True)
            for col in range(COL_SLNO, COL_FIELD_ID + 1):
                ws.cell(row=row, column=col).fill = _SECTION_FILL
            ws.merge_cells(start_row=row, start_column=COL_SLNO, end_row=row, end_column=COL_GUIDANCE)
            row += 1

        sl_no += 1
        ws.cell(row=row, column=COL_SLNO, value=sl_no)
        label_cell = ws.cell(row=row, column=COL_PARAM, value=fld.get("label") or fld["id"])
        label_cell.alignment = Alignment(wrap_text=True, vertical="top")
        if fld.get("required"):
            label_cell.font = Font(bold=True)
        ws.cell(row=row, column=COL_TYPE, value=section.get("title") or "")

        value = answers.get(fld["id"])
        response = ws.cell(row=row, column=COL_RESPONSE, value=value if isinstance(value, str) else "")
        response.protection = Protection(locked=False)
        response.fill = _INPUT_FILL
        response.alignment = Alignment(wrap_text=True, vertical="top")

        guidance = ws.cell(row=row, column=COL_GUIDANCE, value=fld.get("help_text") or "")
        guidance.alignment = Alignment(wrap_text=True, vertical="top")
        ws.cell(row=row, column=COL_FIELD_ID, value=fld["id"])

        _attach_dropdown(ws, fld, row, validations)
        row += 1

    ws.column_dimensions[get_column_letter(COL_SLNO)].width = 8
    ws.column_dimensions[get_column_letter(COL_PARAM)].width = 60
    ws.column_dimensions[get_column_letter(COL_TYPE)].width = 22
    ws.column_dimensions[get_column_letter(COL_RESPONSE)].width = 40
    ws.column_dimensions[get_column_letter(COL_GUIDANCE)].width = 46
    ws.column_dimensions[get_column_letter(COL_FIELD_ID)].hidden = True
    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)
    # Everything except the Response column is locked, so a vendor tabbing
    # through the file cannot shred the questions or the id column.
    ws.protection.sheet = True

    used_titles = {MAIN_SHEET.lower(), META_SHEET.lower()}
    group_sheets = _write_repeatable_groups(wb, schema, answers, used_titles)

    _write_meta_sheet(wb, meta, group_sheets)

    buf = io.BytesIO()
    wb.save(buf)
    wb.close()
    return buf.getvalue()


def _attach_dropdown(ws, fld: dict, row: int, validations: dict[str, DataValidation]) -> None:
    """Add a list validation to a yes_no/select response cell, when it fits.

    Excel caps an inline list formula at 255 characters and cannot express an
    option that itself contains a comma — in either case the cell is simply
    left as free text rather than given a broken dropdown.
    """
    options = fld.get("options")
    if fld.get("type") == "yes_no" and not options:
        options = ["Yes", "No"]
    if not options or fld.get("type") not in ("yes_no", "select"):
        return
    if any("," in str(o) or '"' in str(o) for o in options):
        return
    formula = '"' + ",".join(str(o) for o in options) + '"'
    if len(formula) > 255:
        return
    dv = validations.get(formula)
    if dv is None:
        dv = DataValidation(type="list", formula1=formula, allow_blank=True, showDropDown=False)
        ws.add_data_validation(dv)
        validations[formula] = dv
    dv.add(ws.cell(row=row, column=COL_RESPONSE))


def _write_repeatable_groups(
    wb: Workbook, schema: dict, answers: dict, used_titles: set[str]
) -> list[tuple[str, str]]:
    """One sheet per repeatable group. Returns [(sheet_title, group_id)].

    Row 1 is a hidden row of field ids; row 2 is the human header; rows 3+ are
    the data. Templates produced by form_import.py never have repeatable groups
    today, so in practice this is a no-op that must simply not crash.
    """
    written: list[tuple[str, str]] = []
    for group in schema.get("repeatable_groups") or []:
        fields = group.get("fields") or []
        if not fields:
            continue
        title = _safe_sheet_title(group.get("title") or group["id"], used_titles)
        ws = wb.create_sheet(title)
        for col, fld in enumerate(fields, start=1):
            ws.cell(row=1, column=col, value=fld["id"])
            header = ws.cell(row=2, column=col, value=fld.get("label") or fld["id"])
            header.font = Font(bold=True, color="FFFFFF")
            header.fill = _HEADER_FILL
            ws.column_dimensions[get_column_letter(col)].width = 30

        rows = answers.get(group["id"])
        rows = rows if isinstance(rows, list) else []
        for r, row_data in enumerate(rows, start=3):
            if not isinstance(row_data, dict):
                continue
            for col, fld in enumerate(fields, start=1):
                cell = ws.cell(row=r, column=col, value=row_data.get(fld["id"], ""))
                cell.protection = Protection(locked=False)
        # Leave blank unlocked rows so extra entries can be added.
        for r in range(max(3, len(rows) + 3), len(rows) + 13):
            for col in range(1, len(fields) + 1):
                ws.cell(row=r, column=col).protection = Protection(locked=False)

        ws.row_dimensions[1].hidden = True
        ws.freeze_panes = ws.cell(row=3, column=1)
        ws.protection.sheet = True
        written.append((title, group["id"]))
    return written


def _write_meta_sheet(wb: Workbook, meta: dict[str, Any], group_sheets: list[tuple[str, str]]) -> None:
    ws = wb.create_sheet(META_SHEET)
    rows: list[tuple[str, Any]] = [
        ("format_version", FORMAT_VERSION),
        ("assignment_id", meta.get("assignment_id") or ""),
        ("template_id", meta.get("template_id") or ""),
        ("template_version", meta.get("template_version") or ""),
        ("exported_at", meta.get("exported_at") or ""),
    ]
    rows.extend((f"group_sheet:{title}", group_id) for title, group_id in group_sheets)
    for r, (key, value) in enumerate(rows, start=1):
        ws.cell(row=r, column=1, value=key)
        ws.cell(row=r, column=2, value=str(value))
    ws.sheet_state = "hidden"


# --------------------------------------------------------------------------
# Read
# --------------------------------------------------------------------------


def parse_answer_workbook(data: bytes, schema: dict) -> ParsedAnswers:
    """Read a filled workbook back into an answers dict.

    Cells left blank are omitted entirely rather than written as "", so the
    caller can merge over an existing draft without an untouched cell wiping a
    previously saved answer.
    """
    try:
        wb = load_workbook(filename=io.BytesIO(data), read_only=False, data_only=True)
    except Exception as exc:  # openpyxl raises a grab-bag of types on bad input
        raise ValueError(f"Could not read the workbook: {exc}") from exc

    result = ParsedAnswers()
    try:
        result.meta = _read_meta_sheet(wb)
        if not result.meta:
            result.warnings.append(
                "This file has no GenAIForge Risk metadata sheet, so it could not be verified "
                "against this form. Answers were matched by question text."
            )

        known_ids = {f["id"] for f in _all_fields(schema)}
        by_label = {
            _normalise_label(f.get("label") or ""): f["id"]
            for f in _all_fields(schema)
            if f.get("label")
        }
        options = _option_normaliser(schema)

        if MAIN_SHEET in wb.sheetnames:
            sheet = wb[MAIN_SHEET]
        else:
            sheet = next((s for s in wb.worksheets if s.title != META_SHEET), None)
            if sheet is None:
                raise ValueError("The workbook has no readable sheet")
            result.warnings.append(f'Read answers from sheet "{sheet.title}".')

        _read_main_sheet(sheet, known_ids, by_label, options, result)
        _read_group_sheets(wb, schema, result)
    finally:
        wb.close()

    return result


def _read_meta_sheet(wb) -> dict[str, str]:
    if META_SHEET not in wb.sheetnames:
        return {}
    meta: dict[str, str] = {}
    for row in wb[META_SHEET].iter_rows(min_row=1, max_col=2, values_only=True):
        key = cell_to_str(row[0] if row else None)
        if not key:
            continue
        meta[key] = cell_to_str(row[1] if len(row) > 1 else None)
    return meta


def find_header_row(sheet, scan_rows: int = 10) -> int:
    """Locate the header row. A form description adds a banner above it, so its
    position is not fixed; fall back to the no-description default."""
    for idx, row in enumerate(
        sheet.iter_rows(min_row=1, max_row=scan_rows, max_col=COL_FIELD_ID, values_only=True),
        start=1,
    ):
        if any(cell_to_str(c) == "_field_id" for c in row):
            return idx
    return DEFAULT_HEADER_ROW


def _read_main_sheet(
    sheet,
    known_ids: set[str],
    by_label: dict[str, str],
    options: dict[str, dict[str, str]],
    result: ParsedAnswers,
) -> None:
    header_row = find_header_row(sheet)
    for row in sheet.iter_rows(min_row=header_row + 1, max_col=COL_FIELD_ID, values_only=True):
        cells = list(row) + [None] * (COL_FIELD_ID - len(row))
        field_id = cell_to_str(cells[COL_FIELD_ID - 1])
        label = cell_to_str(cells[COL_PARAM - 1])
        response = cell_to_str(cells[COL_RESPONSE - 1])

        if field_id not in known_ids:
            # Column F stripped or rewritten — fall back to the question text.
            field_id = by_label.get(_normalise_label(label), "")

        if not field_id:
            if response and label:
                result.unmatched.append(label)
            continue
        if not response:
            continue

        canonical = options.get(field_id, {}).get(_normalise_label(response))
        result.answers[field_id] = canonical or response


def _read_group_sheets(wb, schema: dict, result: ParsedAnswers) -> None:
    groups = schema.get("repeatable_groups") or []
    if not groups:
        return
    by_group_id = {g["id"]: g for g in groups}
    # Prefer the sheet->group mapping recorded at export time; fall back to
    # matching the sheet title against the group title.
    mapping = {
        key.split(":", 1)[1]: value
        for key, value in result.meta.items()
        if key.startswith("group_sheet:")
    }
    for sheet in wb.worksheets:
        if sheet.title in (MAIN_SHEET, META_SHEET):
            continue
        group_id = mapping.get(sheet.title)
        if not group_id:
            group_id = next(
                (g["id"] for g in groups if _normalise_label(g.get("title") or "") == _normalise_label(sheet.title)),
                None,
            )
        group = by_group_id.get(group_id or "")
        if not group:
            continue

        header = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), ())
        col_ids = [cell_to_str(c) for c in header]
        valid_ids = {f["id"] for f in group.get("fields") or []}
        if not any(cid in valid_ids for cid in col_ids):
            continue

        rows: list[dict[str, str]] = []
        for raw in sheet.iter_rows(min_row=3, max_col=len(col_ids), values_only=True):
            entry = {
                col_ids[i]: cell_to_str(raw[i])
                for i in range(min(len(col_ids), len(raw)))
                if col_ids[i] in valid_ids and cell_to_str(raw[i])
            }
            if entry:
                rows.append(entry)
        if rows:
            result.answers[group["id"]] = rows


def merge_answers(existing: dict | None, incoming: dict) -> dict:
    """Overlay imported answers onto the current draft — imported wins.

    `incoming` only ever contains non-empty values (parse skips blank cells), so
    a question the vendor left untouched keeps whatever was already in the draft.
    """
    merged = dict(existing or {})
    merged.update(incoming)
    return merged
