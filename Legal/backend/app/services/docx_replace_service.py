"""Apply text replacements inside an existing DOCX while preserving paragraph styles."""

from __future__ import annotations

import io
from typing import Iterable


def apply_replacements_to_docx(
    docx_bytes: bytes,
    replacements: Iterable[tuple[str, str]],
) -> bytes | None:
    """Replace exact substrings in DOCX paragraphs. Returns None if python-docx unavailable."""
    try:
        from docx import Document
    except ImportError:
        return None

    doc = Document(io.BytesIO(docx_bytes))
    changed = False
    for old, new in replacements:
        if not old or old == new:
            continue
        for para in doc.paragraphs:
            if old in para.text:
                _replace_in_paragraph(para, old, new)
                changed = True
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for para in cell.paragraphs:
                        if old in para.text:
                            _replace_in_paragraph(para, old, new)
                            changed = True

    if not changed:
        return docx_bytes

    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def _replace_in_paragraph(para, old: str, new: str) -> None:
    """Replace text in a paragraph, preserving the first run's character style where possible."""
    if old not in para.text:
        return
    full = para.text
    updated = full.replace(old, new, 1)
    if not para.runs:
        para.text = updated
        return
    # Simple path: collapse to one styled run when replacing.
    first = para.runs[0]
    for run in para.runs[1:]:
        run.text = ""
    first.text = updated
