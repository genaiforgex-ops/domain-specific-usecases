"""Parse DOCX into stable, editable parts with deterministic structure hashing."""

from __future__ import annotations

import hashlib
import io
import json
import re
from dataclasses import asdict, dataclass, field
from typing import Any


@dataclass
class DocxPart:
    """A single editable unit inside a DOCX with a stable position-based anchor."""

    id: str
    part_type: str  # paragraph | table_cell | header | footer
    text: str
    path: dict[str, Any]
    style: str | None = None
    section: str = "body"


@dataclass
class DocumentStructure:
    parts: list[DocxPart] = field(default_factory=list)
    structure_hash: str = ""
    full_text: str = ""

    def to_dict(self) -> dict[str, Any]:
        return {
            "parts": [asdict(p) for p in self.parts],
            "structure_hash": self.structure_hash,
            "full_text": self.full_text,
        }

    @classmethod
    def from_dict(cls, data: dict[str, Any]) -> DocumentStructure:
        parts = [DocxPart(**p) for p in data.get("parts", [])]
        return cls(
            parts=parts,
            structure_hash=data.get("structure_hash", ""),
            full_text=data.get("full_text", ""),
        )


def _canonical_payload(parts: list[DocxPart]) -> list[dict[str, Any]]:
    return [
        {
            "id": p.id,
            "part_type": p.part_type,
            "path": p.path,
            "style": p.style,
            "section": p.section,
            "text": p.text,
        }
        for p in parts
    ]


def _compute_hash(parts: list[DocxPart]) -> str:
    payload = json.dumps(_canonical_payload(parts), sort_keys=True, ensure_ascii=False)
    return hashlib.sha256(payload.encode("utf-8")).hexdigest()


def _accepted_text(para) -> str:
    """Paragraph text with tracked changes accepted: insertions (<w:ins>) kept,
    deletions (<w:delText>) dropped. para.text alone misses inserted runs, so a
    redlined doc would yield a structure that ignores AI/human insertions. For a
    document without tracked changes this returns exactly para.text (including
    tabs/breaks), so structure hashes of existing docs are unchanged."""
    from app.services.document_parser import accepted_paragraph_text

    return accepted_paragraph_text(para)


def extract_structure(docx_bytes: bytes) -> DocumentStructure:
    """Extract structured parts from DOCX bytes. Same input yields same structure_hash."""
    from docx import Document

    doc = Document(io.BytesIO(docx_bytes))
    parts: list[DocxPart] = []
    text_blocks: list[str] = []

    for i, para in enumerate(doc.paragraphs):
        text = _accepted_text(para)
        style = para.style.name if para.style else None
        parts.append(
            DocxPart(
                id=f"body:p:{i}",
                part_type="paragraph",
                text=text,
                path={"container": "body", "para_index": i},
                style=style,
                section="body",
            )
        )
        if text.strip():
            text_blocks.append(text)

    for ti, table in enumerate(doc.tables):
        for ri, row in enumerate(table.rows):
            for ci, cell in enumerate(row.cells):
                for pi, para in enumerate(cell.paragraphs):
                    text = _accepted_text(para)
                    style = para.style.name if para.style else None
                    parts.append(
                        DocxPart(
                            id=f"body:tbl:{ti}:r:{ri}:c:{ci}:p:{pi}",
                            part_type="table_cell",
                            text=text,
                            path={
                                "container": "table",
                                "table_index": ti,
                                "row_index": ri,
                                "col_index": ci,
                                "para_index": pi,
                            },
                            style=style,
                            section="body",
                        )
                    )
                    if text.strip():
                        text_blocks.append(text)

    for si, section in enumerate(doc.sections):
        for hf_label, hf in (("header", section.header), ("footer", section.footer)):
            for pi, para in enumerate(hf.paragraphs):
                text = _accepted_text(para)
                style = para.style.name if para.style else None
                parts.append(
                    DocxPart(
                        id=f"{hf_label}:s:{si}:p:{pi}",
                        part_type=hf_label,
                        text=text,
                        path={
                            "container": hf_label,
                            "section_index": si,
                            "para_index": pi,
                        },
                        style=style,
                        section=hf_label,
                    )
                )
                if text.strip():
                    text_blocks.append(text)

    structure_hash = _compute_hash(parts)
    return DocumentStructure(
        parts=parts,
        structure_hash=structure_hash,
        full_text="\n\n".join(text_blocks),
    )


def parts_for_ai(structure: DocumentStructure, max_chars: int = 20_000) -> str:
    """Compact JSON representation for the AI edit planner (always valid JSON)."""
    compact: list[dict[str, str]] = []
    for p in structure.parts:
        if not p.text.strip():
            continue
        entry = {"id": p.id, "type": p.part_type, "text": p.text[:2000]}
        compact.append(entry)
        if len(json.dumps(compact, ensure_ascii=False)) > max_chars:
            compact.pop()
            break
    return json.dumps(compact, ensure_ascii=False)


def part_by_id(structure: DocumentStructure, anchor_id: str) -> DocxPart | None:
    for part in structure.parts:
        if part.id == anchor_id:
            return part
    return None


def last_body_part(structure: DocumentStructure) -> DocxPart | None:
    body_paras = [p for p in structure.parts if p.part_type == "paragraph" and p.section == "body"]
    return body_paras[-1] if body_paras else None


def last_body_part_id_from_dict(structure: dict[str, Any]) -> str | None:
    body_paras = [
        p
        for p in structure.get("parts", [])
        if p.get("part_type") == "paragraph" and p.get("section") == "body"
    ]
    return body_paras[-1]["id"] if body_paras else None


def relevant_text_excerpt(text: str, instruction: str, max_chars: int = 12_000) -> str:
    """Return paragraphs most related to the edit instruction (keeps ADK prompts bounded)."""
    if len(text) <= max_chars:
        return text
    keywords = [w.lower() for w in re.findall(r"[a-zA-Z]{4,}", instruction)]
    paragraphs = [p.strip() for p in text.split("\n\n") if p.strip()]
    if not keywords or not paragraphs:
        return text[:max_chars]

    scored: list[tuple[int, int]] = []
    for i, para in enumerate(paragraphs):
        pl = para.lower()
        score = sum(1 for k in keywords if k in pl)
        if score:
            scored.append((score, i))
    if not scored:
        head = "\n\n".join(paragraphs[:20])
        return head[:max_chars]

    scored.sort(reverse=True)
    indices: set[int] = set()
    for _, i in scored[:25]:
        for j in range(max(0, i - 1), min(len(paragraphs), i + 2)):
            indices.add(j)
    excerpt = "\n\n".join(paragraphs[i] for i in sorted(indices))
    if len(excerpt) > max_chars:
        return excerpt[:max_chars]
    return excerpt
