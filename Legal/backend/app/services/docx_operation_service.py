"""Apply structured DOCX edit operations with high-fidelity preservation."""

from __future__ import annotations

import io
from dataclasses import dataclass, field
from typing import Any

from app.services.docx_replace_service import apply_replacements_to_docx
from app.services.docx_structure_service import DocumentStructure, part_by_id
from app.services.document_parser import (
    content_has_markdown_table,
    insert_content_after_paragraph,
    replace_paragraph_with_content,
)


@dataclass
class OperationApplyResult:
    op_index: int
    op_type: str
    description: str
    status: str  # applied | failed | manual_required
    message: str = ""


@dataclass
class ApplyOperationsResult:
    docx_bytes: bytes | None
    results: list[OperationApplyResult] = field(default_factory=list)
    projected_text: str = ""
    all_applied: bool = False


def project_text_after_operations(
    structure: DocumentStructure,
    operations: list[dict[str, Any]],
) -> str:
    """Simulate operations on part texts for redline preview."""
    texts = {p.id: p.text for p in structure.parts}
    ordered_ids = [p.id for p in structure.parts]

    for op in operations:
        op_type = op.get("op_type", "")
        anchor_id = op.get("anchor_id")
        after_anchor_id = op.get("after_anchor_id")
        content = (op.get("content") or "").strip()
        target_text = op.get("target_text")

        if op_type in ("replace_clause", "replace_span") and anchor_id and anchor_id in texts:
            current = texts[anchor_id]
            if target_text and target_text in current:
                texts[anchor_id] = current.replace(target_text, content, 1)
            elif content:
                texts[anchor_id] = content
        elif op_type == "delete_clause" and anchor_id and anchor_id in texts:
            texts[anchor_id] = ""
        elif op_type in ("insert_clause", "add_definition"):
            insert_id = f"__insert__:{after_anchor_id or anchor_id or 'end'}"
            texts[insert_id] = content
            insert_after = after_anchor_id or anchor_id
            if insert_after and insert_after in ordered_ids:
                idx = ordered_ids.index(insert_after) + 1
                ordered_ids.insert(idx, insert_id)
            else:
                ordered_ids.append(insert_id)

    blocks = [texts[i] for i in ordered_ids if texts.get(i, "").strip()]
    return "\n\n".join(blocks)


def apply_operations_to_docx(
    docx_bytes: bytes,
    structure: DocumentStructure,
    operations: list[dict[str, Any]],
    author: str | None = None,
) -> ApplyOperationsResult:
    """Apply structured operations to DOCX bytes.

    When ``author`` is provided, mutations are emitted as native tracked changes
    (w:ins/w:del) attributed to that author so the edit shows as a redline in
    Word/OnlyOffice. Otherwise text is overwritten in place (legacy behaviour).
    """
    try:
        from docx import Document
        from docx.oxml import OxmlElement
        from docx.text.paragraph import Paragraph
    except ImportError:
        return ApplyOperationsResult(
            docx_bytes=None,
            results=[
                OperationApplyResult(
                    op_index=-1,
                    op_type="",
                    description="",
                    status="failed",
                    message="python-docx not available",
                )
            ],
        )

    if author:
        return _apply_operations_tracked(docx_bytes, structure, operations, author)

    doc = Document(io.BytesIO(docx_bytes))
    results: list[OperationApplyResult] = []
    replacements: list[tuple[str, str]] = []
    direct_ops: list[tuple[int, dict[str, Any]]] = []

    for idx, op in enumerate(operations):
        op_type = op.get("op_type", "")
        description = op.get("description", "")
        anchor_id = op.get("anchor_id")
        after_anchor_id = op.get("after_anchor_id")
        content = (op.get("content") or "").strip()
        target_text = op.get("target_text")

        part = part_by_id(structure, anchor_id) if anchor_id else None
        after_part = part_by_id(structure, after_anchor_id) if after_anchor_id else None

        if op_type in ("replace_clause", "replace_span"):
            if not part:
                results.append(
                    OperationApplyResult(idx, op_type, description, "manual_required", "Anchor not found")
                )
                continue
            old_text = (target_text or part.text or "").strip()
            if not old_text:
                results.append(
                    OperationApplyResult(idx, op_type, description, "manual_required", "No target text")
                )
                continue
            # Tables cannot be applied via run-level string replace — use direct path.
            if content_has_markdown_table(content):
                direct_ops.append((idx, op))
            elif old_text in part.text:
                replacements.append((old_text, content))
                results.append(OperationApplyResult(idx, op_type, description, "applied"))
            else:
                direct_ops.append((idx, op))
        elif op_type == "delete_clause":
            if not part:
                results.append(
                    OperationApplyResult(idx, op_type, description, "manual_required", "Anchor not found")
                )
                continue
            if part.text:
                replacements.append((part.text, ""))
                results.append(OperationApplyResult(idx, op_type, description, "applied"))
            else:
                direct_ops.append((idx, op))
        elif op_type in ("insert_clause", "add_definition"):
            ref_part = after_part or part
            if not ref_part:
                results.append(
                    OperationApplyResult(idx, op_type, description, "manual_required", "Insert anchor not found")
                )
                continue
            direct_ops.append((idx, op))
        else:
            results.append(
                OperationApplyResult(idx, op_type, description, "manual_required", f"Unknown op: {op_type}")
            )

    working_bytes = docx_bytes
    if replacements:
        replaced = apply_replacements_to_docx(working_bytes, replacements)
        if replaced is not None:
            working_bytes = replaced

    doc = Document(io.BytesIO(working_bytes))
    for idx, op in direct_ops:
        op_type = op.get("op_type", "")
        description = op.get("description", "")
        anchor_id = op.get("anchor_id")
        after_anchor_id = op.get("after_anchor_id")
        content = (op.get("content") or "").strip()
        part = part_by_id(structure, anchor_id) if anchor_id else None
        after_part = part_by_id(structure, after_anchor_id) if after_anchor_id else None

        if op_type in ("replace_clause", "replace_span"):
            if not part:
                results.append(
                    OperationApplyResult(idx, op_type, description, "manual_required", "Anchor not found")
                )
                continue
            para = _resolve_paragraph(doc, part.path)
            if para is None:
                results.append(
                    OperationApplyResult(
                        idx, op_type, description, "manual_required", "Exact text not found"
                    )
                )
                continue
            if content_has_markdown_table(content):
                ok = replace_paragraph_with_content(para, content)
            else:
                ok = _set_paragraph_text_preserve_style(para, content)
            if ok:
                results.append(OperationApplyResult(idx, op_type, description, "applied"))
            else:
                results.append(
                    OperationApplyResult(
                        idx, op_type, description, "manual_required", "Exact text not found"
                    )
                )
        elif op_type == "delete_clause":
            if not part:
                results.append(
                    OperationApplyResult(idx, op_type, description, "manual_required", "Anchor not found")
                )
                continue
            para = _resolve_paragraph(doc, part.path)
            if para is not None and _set_paragraph_text_preserve_style(para, ""):
                results.append(OperationApplyResult(idx, op_type, description, "applied"))
            else:
                results.append(
                    OperationApplyResult(idx, op_type, description, "manual_required", "Delete failed")
                )
        elif op_type in ("insert_clause", "add_definition"):
            ref_part = after_part or part
            if not ref_part:
                results.append(
                    OperationApplyResult(idx, op_type, description, "manual_required", "Insert anchor not found")
                )
                continue
            ref_para = _resolve_paragraph(doc, ref_part.path)
            if ref_para is None:
                results.append(
                    OperationApplyResult(idx, op_type, description, "manual_required", "Cannot locate paragraph")
                )
                continue
            if content_has_markdown_table(content):
                insert_content_after_paragraph(ref_para, content)
                results.append(OperationApplyResult(idx, op_type, description, "applied"))
            else:
                new_para = _insert_paragraph_after(ref_para, content, OxmlElement, Paragraph)
                if new_para is not None:
                    results.append(OperationApplyResult(idx, op_type, description, "applied"))
                else:
                    results.append(
                        OperationApplyResult(idx, op_type, description, "manual_required", "Insert failed")
                    )

    out = io.BytesIO()
    doc.save(out)
    current_bytes = out.getvalue()

    projected = project_text_after_operations(structure, operations)
    all_applied = all(r.status == "applied" for r in results) if results else False
    return ApplyOperationsResult(
        docx_bytes=current_bytes,
        results=results,
        projected_text=projected,
        all_applied=all_applied,
    )


def _apply_operations_tracked(
    docx_bytes: bytes,
    structure: DocumentStructure,
    operations: list[dict[str, Any]],
    author: str,
) -> ApplyOperationsResult:
    """Apply operations as native tracked changes attributed to ``author``."""
    from datetime import datetime, timezone

    from docx import Document
    from docx.oxml import OxmlElement
    from docx.text.paragraph import Paragraph

    from app.services.docx_track_changes import (
        delete_paragraph_tracked,
        make_context,
        mark_paragraph_inserted,
        replace_paragraph_tracked,
        replace_span_tracked,
    )

    doc = Document(io.BytesIO(docx_bytes))
    date = datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")
    ctx = make_context(doc, author, date)
    results: list[OperationApplyResult] = []

    # Inserts add paragraphs and shift later body indices; apply them after the
    # index-stable replace/delete ops, and from the bottom of the doc upward.
    edit_ops: list[tuple[int, dict[str, Any]]] = []
    insert_ops: list[tuple[int, dict[str, Any]]] = []
    for idx, op in enumerate(operations):
        if op.get("op_type") in ("insert_clause", "add_definition"):
            insert_ops.append((idx, op))
        else:
            edit_ops.append((idx, op))

    def _body_index(op: dict[str, Any]) -> int:
        ref_id = op.get("after_anchor_id") or op.get("anchor_id")
        ref = part_by_id(structure, ref_id) if ref_id else None
        if ref and ref.path.get("container") == "body":
            return ref.path.get("para_index", 0)
        return -1

    insert_ops.sort(key=lambda item: _body_index(item[1]), reverse=True)

    for idx, op in edit_ops:
        op_type = op.get("op_type", "")
        description = op.get("description", "")
        anchor_id = op.get("anchor_id")
        content = (op.get("content") or "").strip()
        target_text = (op.get("target_text") or "").strip()
        part = part_by_id(structure, anchor_id) if anchor_id else None

        if not part:
            results.append(OperationApplyResult(idx, op_type, description, "manual_required", "Anchor not found"))
            continue
        para = _resolve_paragraph(doc, part.path)
        if para is None:
            results.append(OperationApplyResult(idx, op_type, description, "manual_required", "Cannot locate paragraph"))
            continue

        if op_type in ("replace_clause", "replace_span"):
            if content_has_markdown_table(content):
                # Prefer real Word tables over stuffing GFM pipes into tracked runs.
                ok = replace_paragraph_with_content(para, content)
            elif target_text and target_text in para.text and target_text != content:
                ok = replace_span_tracked(para, target_text, content, ctx)
            else:
                ok = replace_paragraph_tracked(para, content, ctx)
            results.append(
                OperationApplyResult(idx, op_type, description, "applied" if ok else "manual_required", "" if ok else "Exact text not found")
            )
        elif op_type == "delete_clause":
            delete_paragraph_tracked(para, ctx)
            results.append(OperationApplyResult(idx, op_type, description, "applied"))
        else:
            results.append(OperationApplyResult(idx, op_type, description, "manual_required", f"Unknown op: {op_type}"))

    for idx, op in insert_ops:
        op_type = op.get("op_type", "")
        description = op.get("description", "")
        content = (op.get("content") or "").strip()
        ref_id = op.get("after_anchor_id") or op.get("anchor_id")
        ref_part = part_by_id(structure, ref_id) if ref_id else None
        if not ref_part:
            results.append(OperationApplyResult(idx, op_type, description, "manual_required", "Insert anchor not found"))
            continue
        ref_para = _resolve_paragraph(doc, ref_part.path)
        if ref_para is None:
            results.append(OperationApplyResult(idx, op_type, description, "manual_required", "Cannot locate paragraph"))
            continue
        if content_has_markdown_table(content):
            insert_content_after_paragraph(ref_para, content)
            results.append(OperationApplyResult(idx, op_type, description, "applied"))
        else:
            new_para = _insert_paragraph_after(ref_para, content, OxmlElement, Paragraph)
            if new_para is not None:
                mark_paragraph_inserted(new_para, ctx)
                results.append(OperationApplyResult(idx, op_type, description, "applied"))
            else:
                results.append(OperationApplyResult(idx, op_type, description, "manual_required", "Insert failed"))

    out = io.BytesIO()
    doc.save(out)
    results.sort(key=lambda r: r.op_index)
    projected = project_text_after_operations(structure, operations)
    all_applied = all(r.status == "applied" for r in results) if results else False
    return ApplyOperationsResult(
        docx_bytes=out.getvalue(),
        results=results,
        projected_text=projected,
        all_applied=all_applied,
    )


def _resolve_paragraph(doc, path: dict[str, Any]):
    container = path.get("container")
    if container == "body":
        idx = path.get("para_index", 0)
        if 0 <= idx < len(doc.paragraphs):
            return doc.paragraphs[idx]
    elif container == "table":
        ti = path.get("table_index", 0)
        ri = path.get("row_index", 0)
        ci = path.get("col_index", 0)
        pi = path.get("para_index", 0)
        if ti < len(doc.tables):
            table = doc.tables[ti]
            if ri < len(table.rows) and ci < len(table.rows[ri].cells):
                paras = table.rows[ri].cells[ci].paragraphs
                if pi < len(paras):
                    return paras[pi]
    elif container in ("header", "footer"):
        si = path.get("section_index", 0)
        pi = path.get("para_index", 0)
        if si < len(doc.sections):
            section = doc.sections[si]
            hf = section.header if container == "header" else section.footer
            if pi < len(hf.paragraphs):
                return hf.paragraphs[pi]
    return None


def _set_paragraph_text_preserve_style(para, new_text: str) -> bool:
    if not para.runs:
        para.text = new_text
        return True
    first = para.runs[0]
    for run in para.runs[1:]:
        run.text = ""
    first.text = new_text
    return True


def _insert_paragraph_after(ref_para, text: str, OxmlElement, Paragraph):
    new_p = OxmlElement("w:p")
    ref_para._p.addnext(new_p)
    new_para = Paragraph(new_p, ref_para._parent)
    if ref_para.style:
        new_para.style = ref_para.style
    if ref_para.runs:
        new_run = new_para.add_run(text)
        src = ref_para.runs[0]
        new_run.bold = src.bold
        new_run.italic = src.italic
        new_run.underline = src.underline
    else:
        new_para.text = text
    return new_para
