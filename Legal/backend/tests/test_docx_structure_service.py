"""Tests for DOCX structure extraction."""

import io
import unittest

from app.services.docx_structure_service import extract_structure


class DocxStructureTests(unittest.TestCase):
    def test_deterministic_structure_hash(self) -> None:
        try:
            from docx import Document
        except ImportError:
            self.skipTest("python-docx not installed")

        buf = io.BytesIO()
        doc = Document()
        doc.add_paragraph("1. Definitions")
        doc.add_paragraph("Vendor has unlimited liability for all claims.")
        table = doc.add_table(rows=1, cols=2)
        table.rows[0].cells[0].paragraphs[0].add_run("Term")
        table.rows[0].cells[1].paragraphs[0].add_run("Meaning")
        doc.save(buf)
        data = buf.getvalue()

        s1 = extract_structure(data)
        s2 = extract_structure(data)
        self.assertEqual(s1.structure_hash, s2.structure_hash)
        self.assertGreater(len(s1.parts), 0)
        self.assertTrue(any(p.id.startswith("body:p:") for p in s1.parts))
        self.assertTrue(any(p.part_type == "table_cell" for p in s1.parts))

    def test_round_trip_dict(self) -> None:
        try:
            from docx import Document
        except ImportError:
            self.skipTest("python-docx not installed")

        buf = io.BytesIO()
        doc = Document()
        doc.add_paragraph("Sample clause text.")
        doc.save(buf)
        structure = extract_structure(buf.getvalue())
        restored = structure.to_dict()
        self.assertEqual(restored["structure_hash"], structure.structure_hash)
        self.assertIn("Sample", structure.full_text)


if __name__ == "__main__":
    unittest.main()
