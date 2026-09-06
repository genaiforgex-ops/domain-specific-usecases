"""Tests for DOCX operation application."""

import io
import unittest

from app.services.docx_operation_service import apply_operations_to_docx, project_text_after_operations
from app.services.docx_structure_service import extract_structure, last_body_part


class DocxOperationTests(unittest.TestCase):
    def _sample_docx(self) -> bytes:
        from docx import Document

        buf = io.BytesIO()
        doc = Document()
        doc.add_paragraph("Vendor has unlimited liability for claims.")
        doc.add_paragraph("Confidentiality obligations survive termination.")
        doc.save(buf)
        return buf.getvalue()

    def test_replace_span_operation(self) -> None:
        try:
            from docx import Document
        except ImportError:
            self.skipTest("python-docx not installed")

        data = self._sample_docx()
        structure = extract_structure(data)
        liability_part = next(p for p in structure.parts if "unlimited liability" in p.text)
        ops = [
            {
                "op_type": "replace_span",
                "anchor_id": liability_part.id,
                "target_text": "unlimited liability",
                "content": "limited mutual liability",
                "description": "Cap liability",
            }
        ]
        projected = project_text_after_operations(structure, ops)
        self.assertIn("limited mutual liability", projected)
        self.assertNotIn("unlimited liability", projected)

        result = apply_operations_to_docx(data, structure, ops)
        self.assertIsNotNone(result.docx_bytes)
        out = Document(io.BytesIO(result.docx_bytes))
        self.assertIn("limited mutual liability", out.paragraphs[0].text)
        self.assertTrue(any(r.status == "applied" for r in result.results))

    def test_insert_clause_operation(self) -> None:
        try:
            from docx import Document
        except ImportError:
            self.skipTest("python-docx not installed")

        data = self._sample_docx()
        structure = extract_structure(data)
        anchor = last_body_part(structure)
        self.assertIsNotNone(anchor)
        new_clause = "Data protection clause per DPDP Act 2023."
        ops = [
            {
                "op_type": "insert_clause",
                "after_anchor_id": anchor.id,
                "content": new_clause,
                "description": "Add DPDP clause",
            }
        ]
        result = apply_operations_to_docx(data, structure, ops)
        self.assertIsNotNone(result.docx_bytes)
        out = Document(io.BytesIO(result.docx_bytes))
        all_text = "\n".join(p.text for p in out.paragraphs)
        self.assertIn(new_clause, all_text)

    def test_insert_clause_converts_markdown_table_to_word_table(self) -> None:
        try:
            from docx import Document
        except ImportError:
            self.skipTest("python-docx not installed")

        data = self._sample_docx()
        structure = extract_structure(data)
        anchor = last_body_part(structure)
        self.assertIsNotNone(anchor)
        table_md = (
            "Service levels:\n"
            "| Metric | Target |\n"
            "| --- | --- |\n"
            "| Uptime | 99.9% |\n"
            "| Response | 4 hours |\n"
        )
        ops = [
            {
                "op_type": "insert_clause",
                "after_anchor_id": anchor.id,
                "content": table_md,
                "description": "Add SLA table",
            }
        ]
        result = apply_operations_to_docx(data, structure, ops)
        self.assertIsNotNone(result.docx_bytes)
        self.assertTrue(any(r.status == "applied" for r in result.results))
        out = Document(io.BytesIO(result.docx_bytes))
        self.assertGreaterEqual(len(out.tables), 1)
        table = out.tables[0]
        self.assertEqual(table.rows[0].cells[0].text.strip(), "Metric")
        self.assertEqual(table.rows[0].cells[1].text.strip(), "Target")
        self.assertEqual(table.rows[1].cells[0].text.strip(), "Uptime")
        self.assertEqual(table.rows[1].cells[1].text.strip(), "99.9%")
        all_text = "\n".join(p.text for p in out.paragraphs)
        self.assertNotIn("| Metric |", all_text)
        self.assertNotIn("| --- |", all_text)
        self.assertIn("Service levels:", all_text)

    def test_replace_clause_with_markdown_table(self) -> None:
        try:
            from docx import Document
        except ImportError:
            self.skipTest("python-docx not installed")

        data = self._sample_docx()
        structure = extract_structure(data)
        liability_part = next(p for p in structure.parts if "unlimited liability" in p.text)
        table_md = (
            "| Cap | Amount |\n"
            "| --- | --- |\n"
            "| Annual | 12 months fees |\n"
        )
        ops = [
            {
                "op_type": "replace_clause",
                "anchor_id": liability_part.id,
                "target_text": liability_part.text,
                "content": table_md,
                "description": "Replace liability with table",
            }
        ]
        result = apply_operations_to_docx(data, structure, ops)
        self.assertIsNotNone(result.docx_bytes)
        out = Document(io.BytesIO(result.docx_bytes))
        self.assertGreaterEqual(len(out.tables), 1)
        self.assertEqual(out.tables[0].rows[0].cells[0].text.strip(), "Cap")
        body = "\n".join(p.text for p in out.paragraphs)
        self.assertNotIn("| Cap |", body)

    def test_stub_plan_produces_operations(self) -> None:
        try:
            from app.services.ai_service import _stub_plan_docx_operations
        except Exception:
            self.skipTest("app dependencies not installed")

        data = self._sample_docx()
        structure = extract_structure(data)
        plan = _stub_plan_docx_operations(
            structure.to_dict(),
            'replace "unlimited liability" with "limited mutual liability"',
            None,
            None,
            "test-stub",
        )
        self.assertGreater(len(plan.operations), 0)
        self.assertTrue(plan.change_summary)


if __name__ == "__main__":
    unittest.main()
