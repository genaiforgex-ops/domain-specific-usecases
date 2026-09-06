"""Tests for DOCX in-place replacement."""

import io
import unittest

from app.services.docx_replace_service import apply_replacements_to_docx


class DocxReplaceTests(unittest.TestCase):
    def test_replace_in_paragraph(self) -> None:
        try:
            from docx import Document
        except ImportError:
            self.skipTest("python-docx not installed")

        buf = io.BytesIO()
        doc = Document()
        doc.add_paragraph("Vendor has unlimited liability for claims.")
        doc.save(buf)
        original = buf.getvalue()

        updated = apply_replacements_to_docx(
            original, [("unlimited liability", "limited mutual liability")]
        )
        self.assertIsNotNone(updated)
        out = Document(io.BytesIO(updated))
        self.assertIn("limited mutual liability", out.paragraphs[0].text)
        self.assertNotIn("unlimited liability", out.paragraphs[0].text)


if __name__ == "__main__":
    unittest.main()
